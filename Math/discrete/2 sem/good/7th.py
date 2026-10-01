from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import matplotlib.pyplot as plt
import numpy as np

doc = Document()

t = doc.add_heading('Лабораторная работа\nПостроение нечёткой системы Мамдани для выбора коэффициента сжатия видео', 0)
t.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

sections = [
("Введение",
 "Цель работы — разработать нечёткую систему вывода Мамдани для автоматического выбора коэффициента сжатия видео в зависимости от динамичности сцены и критичности размера файла."),
("1. Лингвистические переменные",
 "Входная переменная X — динамичность сцены. Входная переменная Y — критичность размера файла. Выходная переменная Z — коэффициент сжатия."),
]

for h,txt in sections:
    doc.add_heading(h, level=1)
    doc.add_paragraph(txt)

tbl=doc.add_table(rows=1, cols=4)
tbl.style='Table Grid'
for c,txt in zip(tbl.rows[0].cells,["Переменная","Терм 1","Терм 2","Терм 3"]):
    c.text=txt
rows=[("X","Низкая","Средняя","Высокая"),
      ("Y","Слабая","Умеренная","Критическая"),
      ("Z","Мягкое","Оптимальное","Жёсткое")]
for r in rows:
    c=tbl.add_row().cells
    for i,v in enumerate(r): c[i].text=v

doc.add_heading('2. Функции принадлежности', level=1)
doc.add_paragraph('Используются трапециевидные функции для крайних термов и треугольные для средних.')

funcs=[
"Низкая=(0,0,20,40)",
"Средняя=(20,50,80)",
"Высокая=(60,80,100,100)"
]
for f in funcs:
    doc.add_paragraph(f)

x=np.linspace(0,100,1000)
def trap(a,b,c,d):
    y=np.zeros_like(x)
    if b>a:
        m=(x>=a)&(x<b); y[m]=(x[m]-a)/(b-a)
    else:
        y[(x>=a)&(x<=c)]=1
    y[(x>=b)&(x<=c)]=1
    if d>c:
        m=(x>c)&(x<=d); y[m]=(d-x[m])/(d-c)
    return np.clip(y,0,1)
def tri(a,b,c):
    y=np.zeros_like(x)
    m=(x>=a)&(x<=b); y[m]=(x[m]-a)/(b-a)
    m=(x>=b)&(x<=c); y[m]=(c-x[m])/(c-b)
    return np.clip(y,0,1)

plt.figure(figsize=(6,4))
plt.plot(x,trap(0,0,20,40),label='Низкая')
plt.plot(x,tri(20,50,80),label='Средняя')
plt.plot(x,trap(60,80,100,100),label='Высокая')
plt.legend(); plt.grid()
plt.xlabel('Значение'); plt.ylabel('μ')
img1='/mnt/data/fuzzy_graph1.png'
plt.savefig(img1,bbox_inches='tight')
plt.close()

doc.add_picture(img1,width=Inches(5))
doc.add_paragraph('Рисунок 1 – Функции принадлежности.')

doc.add_heading('3. База правил', level=1)
rules=[
("Низкая","Слабая","Оптимальное"),
("Низкая","Умеренная","Жёсткое"),
("Низкая","Критическая","Жёсткое"),
("Средняя","Слабая","Мягкое"),
("Средняя","Умеренная","Оптимальное"),
("Средняя","Критическая","Жёсткое"),
("Высокая","Слабая","Мягкое"),
("Высокая","Умеренная","Мягкое"),
("Высокая","Критическая","Оптимальное"),
]
rt=doc.add_table(rows=1, cols=3)
rt.style='Table Grid'
for c,t in zip(rt.rows[0].cells,['X','Y','Z']): c.text=t
for r in rules:
    row=rt.add_row().cells
    for i,v in enumerate(r): row[i].text=v

doc.add_heading('4. Фаззификация', level=1)
doc.add_paragraph('Для X=40 и Y=75 вычислим степени принадлежности.')
doc.add_paragraph('μX(Средняя)=(40−20)/(50−20)=0.667')
doc.add_paragraph('μY(Умеренная)=(80−75)/(80−50)=0.167')
doc.add_paragraph('μY(Критическая)=(75−60)/(80−60)=0.75')

ft=doc.add_table(rows=1, cols=3)
ft.style='Table Grid'
for c,t in zip(ft.rows[0].cells,['Терм','μ','Комментарий']): c.text=t
vals=[('Средняя(X)','0.667','активна'),
      ('Умеренная(Y)','0.167','активна'),
      ('Критическая(Y)','0.75','активна')]
for r in vals:
    row=ft.add_row().cells
    for i,v in enumerate(r): row[i].text=v

doc.add_heading('5. Агрегирование и активация правил', level=1)
doc.add_paragraph('Используется операция MIN.')
doc.add_paragraph('R1: Средняя ∧ Умеренная → Оптимальное. α1=min(0.667;0.167)=0.167')
doc.add_paragraph('R2: Средняя ∧ Критическая → Жёсткое. α2=min(0.667;0.75)=0.667')

doc.add_heading('6. Аккумулирование', level=1)
doc.add_paragraph('Выходные множества объединяются операцией MAX.')

z=x
opt=np.minimum(tri(20,50,80),0.167)
hard=np.minimum(trap(60,80,100,100),0.667)
agg=np.maximum(opt,hard)

plt.figure(figsize=(6,4))
plt.plot(z,opt,label='Оптимальное (усеч.)')
plt.plot(z,hard,label='Жёсткое (усеч.)')
plt.plot(z,agg,label='MAX')
plt.legend(); plt.grid()
img2='/mnt/data/fuzzy_graph2.png'
plt.savefig(img2,bbox_inches='tight')
plt.close()

doc.add_picture(img2,width=Inches(5))
doc.add_paragraph('Рисунок 2 – Аккумулированное нечёткое множество.')

centroid=np.trapz(z*agg,z)/np.trapz(agg,z)

doc.add_heading('7. Дефаззификация методом центра тяжести', level=1)
doc.add_paragraph('Используется формула:')
doc.add_paragraph('Z* = ∫zμ(z)dz / ∫μ(z)dz')
doc.add_paragraph(f'Численное интегрирование даёт Z* = {centroid:.2f}.')

doc.add_heading('8. Интерпретация результатов', level=1)
doc.add_paragraph(
    f'Полученное значение {centroid:.2f} соответствует достаточно сильному сжатию. '
    'Причина заключается в высокой степени принадлежности терму «Критическая» для переменной размера файла. '
    'Система делает выбор в пользу экономии дискового пространства, однако не переходит к максимально возможному сжатию, '
    'так как динамичность сцены оценивается как средняя.'
)

doc.add_heading('Заключение', level=1)
doc.add_paragraph(
    'В ходе работы была разработана и исследована нечёткая система Мамдани. '
    'Полученная модель может использоваться в программном обеспечении для видеоконвертации и автоматического подбора параметров кодирования.'
)

path='/mnt/data/Academic_Mamdani_Video_Compression.docx'
doc.save(path)
print(path)
